from pathlib import Path
import pymupdf
from docx import Document


def extract_text(file_path: Path) -> dict:
    extension = file_path.suffix.lower()

    if not file_path.is_file():
        raise FileNotFoundError("The stored document was not found.")

    if extension == ".txt":
        text = file_path.read_text(
            encoding="utf-8",
            errors="replace",
        )
        return {
            "text": text,
            "page_count": None,
        }

    if extension == ".pdf":
        pages = []

        with pymupdf.open(file_path) as pdf:
            page_count = len(pdf)

            for page_number, page in enumerate(pdf, start=1):
                page_text = page.get_text("text")
                pages.append(
                    f"[Page {page_number}]\n{page_text}"
                )

        return {
            "text": "\n\n".join(pages),
            "page_count": page_count,
        }


    if extension == ".docx":
        document = Document(file_path)
        parts = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]
                parts.append(" | ".join(cells))

        return {
            "text": "\n".join(parts),
            "page_count": None,
        }

    raise ValueError("Unsupported document format.")
